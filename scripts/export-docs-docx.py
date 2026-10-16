"""Xuất tài liệu dự án ra DOCX. Chạy: python scripts/export-docs-docx.py"""
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

FILES = [
    ("Sổ tay vận hành", "docs/HUONG_DAN_VAN_HANH_HAI_QUAN.md"),
    ("Bàn giao cập nhật theo phản hồi", "docs/HANDOVER_CUSTOMS_FEEDBACK_2026-10-15.md"),
]

NAVY = RGBColor(0x17, 0x32, 0x4D)


def add_md(doc, path):
    with open(path, encoding="utf-8") as f:
        lines = f.read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("### "):
            h = doc.add_heading("", level=3)
            r = h.add_run(line[4:])
        elif line.startswith("## "):
            h = doc.add_heading("", level=2)
            r = h.add_run(line[3:])
        elif line.startswith("# "):
            h = doc.add_heading("", level=1)
            r = h.add_run(line[2:])
            for run in h.runs:
                run.font.color.rgb = NAVY
        elif line.strip().startswith("|") and i + 1 < len(lines) and "---" in lines[i + 1]:
            # bảng markdown
            headers = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(headers))
            table.style = "Light Grid Accent 1"
            for j, htxt in enumerate(headers):
                table.rows[0].cells[j].text = htxt
            for ri, row in enumerate(rows):
                for j in range(min(len(headers), len(row))):
                    table.rows[ri + 1].cells[j].text = re.sub(r"\*\*(.+?)\*\*", r"\1", row[j])
            doc.add_paragraph("")
            continue
        elif re.match(r"^(\d+\.|-|\*) ", line.strip()):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(re.sub(r"^(\d+\.|-|\*) ", "", line.strip()))
        elif line.strip() == "":
            pass
        else:
            p = doc.add_paragraph()
            txt = re.sub(r"\*\*(.+?)\*\*", r"\1", line)
            p.add_run(txt)
        i += 1


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    title = doc.add_heading("TÀI LIỆU DỰ ÁN TALUNG-CUS-WORK", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Hệ thống Quản lý Kế hoạch & Công việc — Chi cục Hải quan khu vực XVI - Hải quan cửa khẩu Quốc tế Tà Lùng")
    for name, path in FILES:
        try:
            add_md(doc, path)
        except FileNotFoundError:
            doc.add_paragraph(f"(Thiếu file {path})")
    out = "docs/Tai-lieu-du-an-Talung.docx"
    doc.save(out)
    print("Đã xuất", out)


main()
